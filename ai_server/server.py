#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
BGTI LMS — AI Auto-Generation Backend Server
============================================
يستقبل ملفات PPTX/PDF من واجهة المدرب، يستخرج نصوصها،
ثم يرسلها لـ Gemini API لتوليد محتوى LMS منظم بصيغة JSON.

تشغيل الخادم:
    cd ai_server
    python server.py

المتطلبات:
    pip install python-pptx pypdf pdfplumber google-generativeai flask flask-cors
"""

import sys
import os
import json
import io
import re
import time

# Force UTF-8 output on Windows
if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

from flask import Flask, request, jsonify
from flask_cors import CORS

# ─── Gemini ───────────────────────────────────────────────────────────────────
import google.generativeai as genai

# ─── File parsers ─────────────────────────────────────────────────────────────
from pptx import Presentation                   # python-pptx
import pypdf                                    # pypdf  (was PyPDF2)
try:
    import pdfplumber                           # optional richer text extraction
    HAS_PDFPLUMBER = True
except ImportError:
    HAS_PDFPLUMBER = False

# ══════════════════════════════════════════════════════════════════════════════
#  CONFIGURATION & .env LOADER
# ══════════════════════════════════════════════════════════════════════════════
def load_dotenv_if_exists():
    env_path = os.path.join(os.path.dirname(__file__), ".env")
    if os.path.exists(env_path):
        with open(env_path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    key, val = line.split("=", 1)
                    key = key.strip()
                    val = val.strip().strip("'\"")
                    if key and val and key not in os.environ:
                        os.environ[key] = val

load_dotenv_if_exists()

GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY", "YOUR_GEMINI_API_KEY_HERE")
GEMINI_MODEL   = os.environ.get("GEMINI_MODEL", "gemini-2.0-flash")      # fast + cheap; switch to gemini-1.5-pro for higher quality
MAX_TEXT_CHARS = 14_000                  # trim extracted text to avoid token limits
PORT           = int(os.environ.get("PORT", 5005))

# ══════════════════════════════════════════════════════════════════════════════
#  FLASK APP
# ══════════════════════════════════════════════════════════════════════════════
app = Flask(__name__)
CORS(app, origins=["*"])   # allow all origins so the local HTML page can call it

# ══════════════════════════════════════════════════════════════════════════════
#  FILE PARSING UTILITIES
# ══════════════════════════════════════════════════════════════════════════════

def extract_text_from_pptx(file_bytes: bytes) -> str:
    """Extract all visible text from a PowerPoint (.pptx) file."""
    prs = Presentation(io.BytesIO(file_bytes))
    lines = []
    for slide_num, slide in enumerate(prs.slides, 1):
        slide_texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_texts.append(shape.text.strip())
        if slide_texts:
            lines.append(f"[Slide {slide_num}]")
            lines.extend(slide_texts)
    return "\n".join(lines)


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from a PDF file using pdfplumber (preferred) or pypdf."""
    if HAS_PDFPLUMBER:
        try:
            with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                pages = []
                for i, page in enumerate(pdf.pages, 1):
                    text = page.extract_text() or ""
                    if text.strip():
                        pages.append(f"[Page {i}]\n{text.strip()}")
            return "\n".join(pages)
        except Exception:
            pass   # fall through to pypdf

    # Fallback: pypdf
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    pages = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[Page {i}]\n{text.strip()}")
    return "\n".join(pages)


def trim_text(text: str, max_chars: int = MAX_TEXT_CHARS) -> str:
    """Trim extracted text to fit within token budget."""
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + "\n\n[... المحتوى مقتطع للحد الأقصى المسموح به ...]"


# ══════════════════════════════════════════════════════════════════════════════
#  GEMINI PROMPT BUILDER
# ══════════════════════════════════════════════════════════════════════════════

LMS_PROMPT_TEMPLATE = """أنت مساعد تعليمي متخصص في تصميم المحتوى التدريبي للغة الإنجليزية التقنية والمصطلحات المهنية.

قم بتحليل المحتوى التعليمي المرفق أدناه وتوليد بيانات مُهيكلة بصيغة JSON باللغتين العربية والإنجليزية.
يجب أن يملأ الناتج المكونات الأربعة لوحدة دراسية أسبوعية في منصة LMS.

## المحتوى المُستخرج من الملف:
---
{extracted_text}
---

## تعليمات الإنتاج:
- استنتج العنوان والمواضيع من المحتوى أعلاه.
- أنشئ 5 إلى 8 مصطلحات إنجليزية-عربية من المصطلحات المذكورة أو ذات الصلة بالمجال.
- أنشئ 5 أسئلة اختيار من متعدد (MCQ) تقيس الفهم.
- أنشئ واجباً تطبيقياً عملياً يستخدم المصطلحات الجديدة.
- الإجابات الصحيحة يجب أن تكون منوعة (ليست كلها A).
- أعد فقط JSON نظيف بدون أي نص إضافي خارج الـ JSON.

## JSON Schema المطلوبة (أعد هذه البنية بالضبط):
{{
  "lecture": {{
    "title": "عنوان المحاضرة المشتق من الملف",
    "duration_minutes": 15,
    "summary": "ملخص شامل ومكثف لمحتوى المحاضرة والنقاط الرئيسية (لا يقل عن 80 كلمة)"
  }},
  "glossary_notes": {{
    "title": "قاموس مصطلحات الأسبوع",
    "description": "وصف القاموس - مثال: أكثر من 6 مصطلحات تقنية مترجمة للعربية",
    "terms": [
      {{"en": "Term 1", "ar": "المصطلح 1", "definition": "شرح مختصر للمصطلح بجملة واحدة"}},
      {{"en": "Term 2", "ar": "المصطلح 2", "definition": "شرح مختصر للمصطلح بجملة واحدة"}}
    ]
  }},
  "quiz": {{
    "title": "اختبار الأسبوع: [اسم الموضوع]",
    "duration_minutes": 15,
    "total_grade": 10,
    "questions": [
      {{
        "type": "mcq",
        "question": "نص السؤال بالإنجليزية أو العربية",
        "options": ["الخيار أ", "الخيار ب", "الخيار ج", "الخيار د"],
        "correct_answer": "الخيار أ"
      }}
    ]
  }},
  "assignment": {{
    "title": "الواجب الأسبوعي: [عنوان التكليف]",
    "total_grade": 10,
    "instructions": "تعليمات الواجب التطبيقي - مثال: استخدم المصطلحات الجديدة في 5 جمل كاملة",
    "questions": [
      {{
        "question": "سؤال تطبيقي أو اختيار من متعدد للواجب",
        "options": ["الخيار أ", "الخيار ب", "الخيار ج", "الخيار د"],
        "correct_answer": "الخيار أ"
      }}
    ]
  }}
}}"""


def build_lms_content(extracted_text: str) -> dict:
    """Send extracted text to Gemini and parse the JSON response."""
    if GEMINI_API_KEY == "YOUR_GEMINI_API_KEY_HERE":
        raise ValueError("لم يتم تعيين GEMINI_API_KEY. يرجى إضافة مفتاح API الخاص بك.")

    genai.configure(api_key=GEMINI_API_KEY)
    model = genai.GenerativeModel(
        model_name=GEMINI_MODEL,
        generation_config=genai.GenerationConfig(
            temperature=0.4,
            response_mime_type="application/json",
        ),
    )

    prompt = LMS_PROMPT_TEMPLATE.format(extracted_text=trim_text(extracted_text))
    response = model.generate_content(prompt)
    raw = response.text.strip()

    # Strip markdown code fences if the model wraps output in ```json ... ```
    raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
    raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
    raw = raw.strip()

    return json.loads(raw)


def build_demo_content_from_text(extracted_text: str, filename: str) -> dict:
    """
    Fallback extractor: Parses lines from the extracted text to create a structured
    LMS module following the exact requested JSON schema.
    Used when GEMINI_API_KEY is not configured yet or during offline development.
    """
    lines = [l.strip() for l in extracted_text.splitlines() if l.strip() and not l.startswith("[")]
    title = lines[0] if lines else os.path.splitext(filename)[0]
    
    # Extract terms (looking for lines with colons or bullet patterns)
    extracted_terms = []
    for line in lines[1:]:
        if ":" in line:
            parts = line.split(":", 1)
            term_en = re.sub(r"^\d+[\.\)]\s*", "", parts[0]).strip()
            desc = parts[1].strip()
            extracted_terms.append({
                "en": term_en,
                "ar": f"مصطلح: {term_en}",
                "definition": desc or "مصطلح تقني تخصصي معتمد في بيئة العمل الصناعية والتدريبية"
            })
        if len(extracted_terms) >= 6:
            break

    if not extracted_terms:
        clean_title = title.replace("دورة", "").replace("مقرر", "").strip() or "المحتوى التخصصي"
        extracted_terms = [
            {"en": f"{clean_title} Core Concept", "ar": f"المفهوم الأساسي لـ {clean_title}", "definition": f"الركيزة النظرية والتطبيقية المعتمدة في دراسة وممارسة {clean_title}."},
            {"en": f"{clean_title} Practical Workflow", "ar": f"مسار العمل التطبيقي لـ {clean_title}", "definition": f"تسلسل الخطوات التنفيذية المنهجية المتبعة لتطبيق مهارات {clean_title} بدقة."},
            {"en": f"{clean_title} Standards", "ar": f"المعايير المعتمدة لـ {clean_title}", "definition": f"المواصفات الفنية والمهنية التي تضمن جودة المخرجات في {clean_title}."},
            {"en": f"{clean_title} Verification", "ar": f"التحقق وضبط جودة {clean_title}", "definition": f"إجراءات الفحص والتقييم الدوري لضمان مطابقة مخرجات {clean_title} للمواصفات."}
        ]

    summary_text = " ".join(lines[:6]) if len(lines) >= 3 else extracted_text[:400]

    return {
        "lecture": {
            "title": f"محاضرة: {title}",
            "duration_minutes": 15,
            "summary": f"تتناول هذه المحاضرة المفاهيم الأساسية المستخلصة من الملف التدريبي، مع التركيز على المصطلحات المهنية والتطبيق العملي في المنشآت الصناعية. المحاور الأساسية: {summary_text[:350]}..."
        },
        "glossary_notes": {
            "title": f"قاموس مصطلحات الأسبوع: {title}",
            "description": f"أكثر من {len(extracted_terms)} مصطلحات فنية إنجليزية مترجمة للعربية وموضحة وفق المنهج التدريبي المعتمد.",
            "terms": extracted_terms
        },
        "quiz": {
            "title": f"اختبار الأسبوع: {title}",
            "duration_minutes": 15,
            "total_grade": 10,
            "questions": [
                {
                    "type": "mcq",
                    "question": f"What is the primary definition of '{extracted_terms[0]['en']}' according to the technical lecture?",
                    "options": [
                        extracted_terms[0]["definition"][:60],
                        "An unrelated manual hand tool",
                        "A temporary administrative report",
                        "None of the above"
                    ],
                    "correct_answer": extracted_terms[0]["definition"][:60]
                },
                {
                    "type": "mcq",
                    "question": f"ما هو المقابل الدقيق للمصطلح المهني '{extracted_terms[min(1, len(extracted_terms)-1)]['en']}'؟",
                    "options": [
                        extracted_terms[min(1, len(extracted_terms)-1)]["ar"],
                        "أداة تشغيل عامة",
                        "مستند محاسبي دوري",
                        "إجراء صيانة طارئة"
                    ],
                    "correct_answer": extracted_terms[min(1, len(extracted_terms)-1)]["ar"]
                },
                {
                    "type": "mcq",
                    "question": f"Which term best describes: '{extracted_terms[min(2, len(extracted_terms)-1)]['definition'][:70]}'?",
                    "options": [
                        extracted_terms[min(2, len(extracted_terms)-1)]["en"],
                        "General Tool",
                        "Standard Blueprint",
                        "Safety Barrier"
                    ],
                    "correct_answer": extracted_terms[min(2, len(extracted_terms)-1)]["en"]
                },
                {
                    "type": "mcq",
                    "question": "ما الغرض الرئيسي من دراسة المصطلحات الفنية باللغة الإنجليزية في بيئة العمل؟",
                    "options": [
                        "التواصل الدقيق وفهم كتيبات التشغيل والرموز الهندسية العالمية",
                        "إلغاء استخدام اللغة العربية في المحادثات اليومية",
                        "حفظ الكلمات دون تطبيقها في الورش والمختبرات",
                        "إجراء اختبارات نظرية فقط"
                    ],
                    "correct_answer": "التواصل الدقيق وفهم كتيبات التشغيل والرموز الهندسية العالمية"
                },
                {
                    "type": "mcq",
                    "question": f"True or False: Technical terminology in '{title}' requires adherence to international safety standards.",
                    "options": [
                        "True (صواب - إلزامي في المعايير الصناعية)",
                        "False (خطأ - غير مطلوب)",
                        "Optional in minor tasks",
                        "Only applies to management"
                    ],
                    "correct_answer": "True (صواب - إلزامي في المعايير الصناعية)"
                }
            ]
        },
        "assignment": {
            "title": f"الواجب الأسبوعي: تطبيق مصطلحات {title}",
            "total_grade": 10,
            "instructions": f"المطلوب من المتدرب: 1) استخدام 5 مصطلحات من قاموس هذا الأسبوع في جمل مهنية تصف مهمة في ورشة التدريب. 2) كتابة تقرير موجز لا يتجاوز 100 كلمة يوضح فيه كيفية تطبيق إجراءات السلامة المذكورة في المحاضرة.",
            "questions": [
                {
                    "question": f"اختر المصطلح الأنسب لإكمال الجملة الفنية التالية: 'Before starting the operation, the technician must inspect the _______'.",
                    "options": [
                        extracted_terms[0]["en"],
                        "Office chair",
                        "Coffee cup",
                        "Calendar"
                    ],
                    "correct_answer": extracted_terms[0]["en"]
                },
                {
                    "question": f"كيف تترجم العبارة التالية بدقة مهنية: 'Always wear the required safety gear on site'؟",
                    "options": [
                        "ارتدِ دائماً معدات السلامة المطلوبة في موقع العمل",
                        "احتفظ بأدوات الصيانة في المستودع",
                        "سجل الحضور اليومي في البوابة",
                        "أوقف تشغيل المعدات عند نهاية الدوام"
                    ],
                    "correct_answer": "ارتدِ دائماً معدات السلامة المطلوبة في موقع العمل"
                }
            ]
        }
    }


# ══════════════════════════════════════════════════════════════════════════════
#  ROUTES
# ══════════════════════════════════════════════════════════════════════════════

@app.route("/health", methods=["GET"])
def health():
    """Quick health-check endpoint."""
    return jsonify({"status": "ok", "model": GEMINI_MODEL, "port": PORT})


@app.route("/api/ai-generate", methods=["POST", "OPTIONS"])
@app.route("/api/generate", methods=["POST", "OPTIONS"])
def ai_generate():
    """
    Accepts multipart/form-data or application/json (with base64)
    (.pptx or .pdf), extracts text, calls Gemini, returns structured JSON.
    """
    if request.method == "OPTIONS":
        return jsonify({"status": "ok"}), 200

    t0 = time.time()

    # Support application/json with base64 (matching Vercel serverless interface)
    file_bytes = b""
    filename = ""
    if request.is_json:
        data = request.get_json() or {}
        filename = data.get("filename", "upload.pptx")
        b64 = data.get("base64", "")
        if b64:
            import base64 as b64module
            try:
                file_bytes = b64module.b64decode(b64)
            except Exception as e:
                return jsonify({"error": f"فشل فك تشفير Base64: {e}"}), 400
    elif "file" in request.files:
        uploaded = request.files["file"]
        filename = uploaded.filename or ""
        file_bytes = uploaded.read()
    else:
        return jsonify({"error": "لم يتم استلام أي ملف. يرجى رفع ملف PPTX أو PDF."}), 400

    ext = os.path.splitext(filename)[1].lower()
    if ext not in (".pptx", ".pdf"):
        return jsonify({
            "error": f"نوع الملف غير مدعوم: '{ext}'. الأنواع المدعومة: .pptx و .pdf فقط."
        }), 415

    if len(file_bytes) == 0:
        return jsonify({"error": "الملف المرفوع فارغ. يرجى رفع ملف يحتوي على محتوى."}), 400

    # ── 2. Extract text ───────────────────────────────────────────────────────
    try:
        if ext == ".pptx":
            extracted_text = extract_text_from_pptx(file_bytes)
        else:
            extracted_text = extract_text_from_pdf(file_bytes)
    except Exception as ex:
        return jsonify({"error": f"فشل في قراءة الملف: {str(ex)}"}), 422

    if not extracted_text.strip():
        return jsonify({
            "error": "لم يتم العثور على نص قابل للقراءة في الملف. قد يكون الملف محمياً أو يحتوي على صور فقط."
        }), 422

    # ── 3. Call Gemini (or fallback demo generator if key not set) ────────────
    is_demo = False
    try:
        if GEMINI_API_KEY == "YOUR_GEMINI_API_KEY_HERE" or not GEMINI_API_KEY:
            is_demo = True
            lms_data = build_demo_content_from_text(extracted_text, filename)
        else:
            lms_data = build_lms_content(extracted_text)
    except ValueError as ve:
        return jsonify({"error": str(ve)}), 503
    except json.JSONDecodeError as je:
        return jsonify({"error": f"الذكاء الاصطناعي أعاد بيانات غير صالحة. يرجى المحاولة مرة أخرى. ({je})"}), 502
    except Exception as ex:
        # If API call fails (e.g. quota, bad key), fall back gracefully
        print(f"[!] Gemini call failed: {ex}. Falling back to smart text-based extractor.")
        is_demo = True
        lms_data = build_demo_content_from_text(extracted_text, filename)

    elapsed = round(time.time() - t0, 2)

    return jsonify({
        "success": True,
        "is_demo": is_demo,
        "elapsed_seconds": elapsed,
        "filename": filename,
        "extracted_chars": len(extracted_text),
        "data": lms_data
    })


# ══════════════════════════════════════════════════════════════════════════════
#  ENTRY POINT
# ══════════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    print("=" * 60)
    print(" BGTI LMS — AI Auto-Generation Server")
    print(f" Model  : {GEMINI_MODEL}")
    print(f" Port   : {PORT}")
    print(f" API Key: {'✓ Set' if GEMINI_API_KEY != 'YOUR_GEMINI_API_KEY_HERE' else '✗ NOT SET — edit server.py'}")
    print("=" * 60)
    app.run(host="0.0.0.0", port=PORT, debug=False)
